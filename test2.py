"""temp=49
x=[]
i=2
	
while temp!=0:
	if temp%i==0:
		temp = temp//i
		x.append(i)
	else:
		i+=1"""

#print(x)
"""a=[1,"Aman",2.3,4]
for i in range(len(a)):
	print(a[i])"""
"""a=[1,"Aman",2.3,4]
for i in a:
	print(i)

a=
b=[]
for i in range(a):
	b.append(a[i])
print(b)"""

"""x = abs(3+5j)
print(type(x))"""


"""def fact(n):
 	x = []
 	for i in range(2,n):
 		if n%i == 0 :
 			x.append(i)
 	return x 
 
 def prime(y):
 	z = []
 	for j in y :
 		for k in range(2,j):
 			if j%k == 0 :
 				break
 		else : 
 			z.append(j)
 	return z
 
 def multiple_occurance(n,Z):
 	global result 
 	temp = n
 	count = 0
 	for m in Z : 
 		while temp >= 1 and temp % m == 0 :
 			temp = temp // m 
 			count = count + 1 
		result[m] = count
		count = 0

n = int(input("Enter Value : "))
result = dict()

y = fact(n)
Z = prime(y)
print(y)
print(Z)
multiple_occurance(n,Z)
print(result)

for k in (result.values()):
	print(result.keys(),end="")
	if result.values()!=1:	
		print("*",end="")"""

"""user=[2,7,11,15]
target=9
for i in user:

	i+i==target
	print(i)"""

"""var1=int(input("Enter number:-"))
reverse = 0
var2=[]
num=0
while var1!=0 and num<9:

	digit = var1%10
	var1//=10
	reverse=reverse*10+digit
	digit=reverse%10
	var2.append(digit)
	num=digit+num
		#while len(num) !=index:


print(num)"""



"""num=[]
for i in range (len(var2)):
	n=var2[-1]+var2[-2]
	num.append(n)

print(num)"""

"""def sum_of_digits(num):
  	total=0
  	while num>0:
		total += n % 10  
		n //= 10
	return t
number =584447788
result = sum_of_digits(number)
print(result)"""
"""
def sum_until_value(arr, limit): 
    total = 0 
    for num in arr: 
        if total + num > limit: 
            break 
        total += num 
    return total """


"""arr = [1, 2, 3, 4, 5] 
limit = 7 
total = [] 
for i in range (len(arr)): 
	for j in range (i+1,len(arr)):
		if arr[i] + arr[j] ==limit:
			total.append(i)
			total.append(j)
			print(total)"""


"""def sum_of_digits(num):
  	total=0
  	fact=True
  	while fact==True:
  		if num>0 :
  			total+= num%10
  			num //= 10
  			if total>9:
  				total=sum_of_digits(total)		
  		else:
  			fact=False
  			return total
  			  		
number =12365498778966574159
result = sum_of_digits(number)
print(result)"""
		

"""list1=["a.json","b.json","b.text","p.py"]
list2=[]
for i in range (len(list1)):
	list1[i]=list1[i].split(".")[0]
print(list1)"""

"""l1=[2,4,3]
l2=[5,6,4]
l3=[]
var1=''
var2=""
for i in range (len(l1)):
	var1+=var1.join(str(l1[i]))
#print(var1)
for j in range (len(l2)):
	var2+=var2.join(str(l2[j]))
		
#print(var2)
print(int(var1)+int(var2))"""

"""def addTwoNumbers(l1,l2):
    var1=''
    var2=""
    var=None
    for i in range (len(l1)):
    	var1+=var1.join(str(l1[i]))
    for j in range (len(l2)):
    	var2+=var2.join(str(l2[j]))
    var=int(var1)+int(var2)
    return var

l1=[2,4,3]
l2=[5,6,4]
print(addTwoNumbers(l1,l2))"""
"""s="Hello world?#"
print(s[0:5])"""
"""for i in range (1,10):
	print(i)
	i+=1
print(i)"""

"""n="nitesh"
for i in n:
	x=i.upper()	
	print(x)"""


"""def reverse(x):

	new=0
	var=abs(x)	
	while var!=0 and var>0:
		digit=var%10
		var//=10
		#digit=str(digit)
		new=new*10+digit
	if x>0:
		return new
	else:
		return (-new)


x=-120145
print(reverse(x))"""

"""def reverse(x):
	new=0
	var=abs(x)
	if -2**31<= x <=2**31-1 :
		return 0
	else:		
		while var!=0 and var>0:
			digit=var%10
			var//=10
			new=new*10+digit
		if x>0:
			return new
		else:
			return (-new)


x=-21474836482

print(reverse(x))"""


"""class squre:
	def __init__(self):
		self.side=int(input("enter side::"))

	def area(self):
		return self.side*self.side

	def perimeter(self):
		return 4*self.side
s1=squre()
p1=s1.perimeter()
a1=s1.area() 
print("Area of squre with side {}={}".format(s1.side,a1))"""




"""def numb(number): 	
	#number=input("Enter number::")	
	n=number.isnumeric()
	if len(str(number))==10 and n is True and len(str(number[0]))!=0:
		return number
	else:
		print("=======Please Enter valid Number=======")
		print("=== The number will not starte on 0 ===")
		print("== All Digits are numeric are compulsory ==")		
		numb()
num=input("enter number::")
numb(num)"""


"""def numb(n): 	
	#number=input("Enter number::")	
	n=number.isnumeric()

	if len(str(number))==10 and n is True and list(number[0])==0:
		print( number)
	else:
		print("=======Please Enter valid Number=======")
		print("=== The number will not starte on 0 ===")
		print("== All Digits are numeric compulsory ==")		
		numb(n)
	break



number=input("entre number::")
numb(number)"""
#===================##valid and invalid##=========================
"""string='[{()}[]]'
string=list(string)
open1=["[","{","("]
close1=["]","}",")"]
count=0
for i in string:	
	if i in open1:
		count+=1
	elif i in close1:
		count-=1
if count==0:
	print("valid")
elif count>0 and count<0:	
	print("Invalid")"""

#================================================================
"""a=int(input("enter number:"))
x=[]
for i in range (1,a+1):
	b=int(input("enter:"))
	x.append(b)
print(x)
for i in x:
	if i <=5:
		print(i)"""

#=======================================================
"""a=(int(input()),int(input()),int(input()))
if a[0]> a[1]:
	print(a[0])
else:
	print(a[1])"""

	# import openpyxl module
"""import openpyxl

# import BarChart class from openpyxl.chart sub_module
from openpyxl.chart import BarChart,Reference

# Call a Workbook() function of openpyxl 
# to create a new blank Workbook object
wb = openpyxl.Workbook()

# Get workbook active sheet 
# from the active attribute.
sheet = wb.active

# write o to 9 in 1st column of the active sheet
for i in range(10):
	sheet.append([i])

# create data for plotting
values = Reference(sheet, min_col = 1, min_row = 1,
						max_col = 1, max_row = 10)

# Create object of BarChart class
chart = BarChart()

# adding data to the Bar chart object
chart.add_data(values)

# set the title of the chart
chart.title = " BAR-CHART "

# set the title of the x-axis
chart.x_axis.title = " X_AXIS "

# set the title of the y-axis
chart.y_axis.title = " Y_AXIS "

# add chart to the sheet
# the top-left corner of a chart
# is anchored to cell E2 .
sheet.add_chart(chart, "E2")

# save the file
wb.save("barChart.xlsx")"""



"""class stack:
    #docstring for stack
    def __init__(self):
        self.stacklist  = []
        self.max = -1 

    def show(self) : 
        print("Stacklist : ", self.stacklist)

    def push(self,x):
        self.stacklist.append(x)
        self.max = self.max + 1 

    def pop(self) : 
        z = self.stacklist.pop()
        self.max = self.max-1
        return z 
        

operator={"/":2,"*":2,"-":1,"+":1}
post_exp=["24","4","/","78","/","9","*","8","2","/","6","*","+","9","+","2","-"]
result= stack()
new=[]
for i in post_exp:
	int(i)
	new.append(i)
print(new)

for i in post_exp:
    if i not in operator:
    	result.push(i)
    else:
    	x=result.pop()
    	y=result.pop()
    	z=int(x)++int(y)
    	result.push(z)
else:
	result.append()
print(Temp)"""



        
